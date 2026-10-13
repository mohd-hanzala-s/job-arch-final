"""Checks for IRIS Job Architecture Model v2.1 · five-step weightage and ratings, live re-scoring. Exit 0 = all pass. Usage: python3 verify_model.py"""
import collections
import os
import re
import shutil
import sys
import tempfile
import zipfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ja_common as J  # noqa: E402
import ja_raw as RW  # noqa: E402
import ja_scale as SC  # noqa: E402
import ja_v114 as V  # noqa: E402
import ja_xlsx as X  # noqa: E402
import build_model as BM  # noqa: E402
from build_model import EXP_RANGE  # noqa: E402

OUTFILE = os.environ.get("JA_FILE", BM.MASTER_FILE)
SIMPLE = "Simplified" in os.path.basename(OUTFILE)
N = dict(zip(("HOW", "LAD", "RL", "RM", "RJ", "MB", "TM", "PC", "RAW", "BE", "REF"), BM.NAMES[SIMPLE]))
from build_bh import PARAMS, PKEYS  # noqa: E402

RES = []


def check(name, ok, detail=""):
    RES.append(bool(ok))
    print(("PASS " if ok else "FAIL ") + name + (f" — {detail}" if detail else ""))


def main():
    wb = openpyxl.load_workbook(OUTFILE, read_only=True, data_only=True)
    want = [n for n in BM.NAMES[SIMPLE] if not n.startswith("zz")]
    check("Sheets in the order of the story, References & Benchmarks last", wb.sheetnames == want, str(wb.sheetnames))
    t = list(wb[N["TM"]].iter_rows(min_row=6, values_only=True))
    h = {x: i for i, x in enumerate(t[0])}
    d = [r for r in t[1:] if r[h["Emp ID"]] or r[h["Employee name"]]]
    check("Every person, top down from the CEO", len(d) == 586 and d[0][h["Employee name"]] == "Balachandran K", f"{len(d)}")
    boiler = sum(1 for ws in wb.worksheets if ws.title != N["REF"] for row in ws.iter_rows(values_only=True) for v in row
                 if isinstance(v, str) and re.search(r"no survey answers or KRAs on file|rated like \d+ colleagues|kept at the start point|KRAs not received|No survey response|not in the files received", v))
    ph = sum(1 for ws in wb.worksheets if ws.title in (N["TM"], N["RAW"], N["PC"]) for row in ws.iter_rows(min_row=7, values_only=True)
             for v in row if isinstance(v, str) and not v.startswith("=") and RW.PLACEHOLDER.match(v))
    check("No-input cells are blank — no placeholder text anywhere", boiler == 0 and ph == 0, f"{boiler} boilerplate · {ph} 'not provided'-type cells")
    check("Leader columns blank until a leader gives an input", all(r[h["Leader proposed level"]] in (None, "") for r in d))
    mc = collections.Counter(r[h["Manager check · live"]] for r in d)
    flagged = sum(1 for r in d if r[h["Structure flag"]])
    exc = mc.get("Structure exception · chain deeper than the bands", 0)
    check("No overlap: every reportee sits in a lower band than their manager — CEO and Chief Officers in Band E agreed; a chain "
          "deeper than the seven bands is the only marked structure exception, still one level below the manager",
          not any(str(k).startswith("⚠") for k in mc) and exc == flagged and flagged <= 1, f"{dict(mc)} · flagged {flagged}")
    orig = all(J.text(r[h["Reporting manager"]]).find(" · was ") < 0 for r in d)
    check("Everyone reports to their original manager", orig)
    check("Before any input, the level in force is IRIS's level for everyone and the band rule holds nobody back",
          all(r[h["Level in force · after the band rule"]] == r[h["IRIS proposed level"]] and not r[h["Band rule · live"]] for r in d))
    regular = [lv for lv in V.LEVELS_STD if lv not in ("A1", "E2")]
    overl = [(a, b) for a, b in zip(regular, regular[1:]) if EXP_RANGE[a][1] > EXP_RANGE[b][0]]
    check("Ladder experience ranges do not overlap", not overl, str(overl))
    seen = collections.defaultdict(set)
    for r in d:
        if r[h["Job family"]] in (V.HEAD_FAMILY, V.ENT_FAMILY):
            continue
        seen[(r[h["Role ID"]], r[h["Lane"]], J.text(r[h["Standard title for the role and level"]]).replace(" – POD Lead", ""))].add(r[h["IRIS proposed level"]])
    dup = [k for k, v in seen.items() if len(v) > 1]
    check("One standard title per role and level", not dup, f"{len(dup)} {dup[:3]}")
    if not SIMPLE:
        rm = list(wb[N["RM"]].iter_rows(min_row=6, values_only=True))
        rh = {x: i for i, x in enumerate(rm[0])}
        role_rows = [r for r in rm[1:] if r[rh["Row"]] == "ROLE"]
        ids = {r[rh["Role ID"]] for r in role_rows}
        check("Every person's role is in the Role Definitions", all(r[h["Role ID"]] in ids for r in d), f"{len(ids)} roles")
        vac_def = [r for r in role_rows if r[rh["People"]] == "Vacant"]
        check("Role Definitions include roles nobody holds yet, with every level", len(vac_def) >= 5 and
              all(sum(1 for x in rm[1:] if x[rh["Role ID"]] == r[rh["Role ID"]] and x[rh["Row"]] == "level") >= 2 for r in vac_def),
              f"{len(vac_def)} vacant roles")
        check("Role Definitions: one block per role (sub-family), every level with its IC / Manager titles",
              len(role_rows) == len({(r[rh["Role · sub-family"]], r[rh["Job family"]]) for r in role_rows}))
    rl = list(wb[N["RL"]].iter_rows(min_row=6, values_only=True))
    lh = {x: i for i, x in enumerate(rl[0])}
    lrows = [r for r in rl[1:] if r[lh["Row"]] == "level"]
    keys = {r[lh["Key"]] for r in rl[1:]}
    check("Every placed person reads their weightage from their own role × level × lane row",
          all(r[h["Ladder key"]] in keys and r[h["Ladder key"]].split("|")[1] == r[h["IRIS proposed level"]]
              and r[h["Ladder key"]].split("|")[0] == r[h["Role ID"]] for r in d
              if V.band(r[h["IRIS proposed level"]]) != "X"), f"{len(lrows)} level rows")
    role_rows_l = [r for r in rl[1:] if r[lh["Row"]] == "ROLE"]
    gen = [r for r in role_rows_l if r[lh["Job family"]] not in (V.HEAD_FAMILY, V.ENT_FAMILY)]
    lanes_at = collections.defaultdict(set)
    for r in lrows:
        lanes_at[r[lh["Role ID"]]].add((r[lh["Level"]], r[lh["Lane"]]))
    want_rows = {(lv, ln) for lv in V.LEVELS_PC for ln in (("IC", "Manager") if V.band(lv) in ("B", "C") else ("All",))
                 if not (ln == "All" and V.VAL[lv] > V.VAL["D1"] and False)}
    short_ = [r[lh["Role ID"]] for r in gen if not want_rows <= lanes_at[r[lh["Role ID"]]]]
    check("A role is a sub-family: every role runs A1 to E2, with IC and Manager lanes in Bands B and C — vacant levels and lanes included",
          not short_ and len(role_rows_l) == len({r[lh["Role · sub-family"]] + r[lh["Job family"]] for r in role_rows_l}),
          f"{len(role_rows_l)} roles · {sum(1 for r in lrows if r[lh['People in IRIS']] == 'Vacant')} vacant rows · {short_[:3]}")
    check("Roles nobody holds yet have a ladder too", sum(1 for r in role_rows_l if r[lh["People in IRIS"]] == "Vacant") >= 5,
          str(sum(1 for r in role_rows_l if r[lh["People in IRIS"]] == "Vacant")))
    wcols = [lh[x] for x in lh if str(x).startswith(("IRIS ·", "Market ·", "In use ·"))]
    codes = {c for c, *_ in SC.W5} | {0}
    check("Every level row has IRIS, Market and In-use weightage in five steps (0 = not weighed) and a written justification",
          all(r[c] in codes for r in lrows for c in wcols) and
          all("MARKET —" in J.text(r[lh["Justification · IRIS and Market"]]) and "IRIS —" in J.text(r[lh["Justification · IRIS and Market"]]) for r in lrows))
    scols = [lh[x] for x in lh if str(x).startswith("Share ·")]
    bad = [r[lh["Key"]] for r in lrows if abs(sum(float(r[c] or 0) for c in scols) - 1) > 1e-6]
    check("Each level row's shares of the decision add up to 100%", not bad, f"{len(bad)} rows off {bad[:2]}")
    check("Nothing changed yet: no level row reports a change", all((r[lh["Changes here"]] or 0) == 0 for r in rl[1:]))
    lt = collections.defaultdict(set)
    for r in lrows:
        lt[(r[lh["Role ID"]], r[lh["Lane"]], r[lh["Standard title at this level"]])].add(r[lh["Level"]])
    check("One standard title per role and level on the ladder, vacant levels included", all(len(v) == 1 for v in lt.values()))
    sv = [r[h["Inputs received in the survey"]] for r in d]
    check("Inputs received in the survey: a detailed summary for everyone who answered, blank for the rest",
          sum(1 for x in sv if x) >= 500 and all((x or "").startswith(("Answered", "Shared")) for x in sv if x)
          and all(not x or r[h["Survey · questions answered"]] or "CARBON" in x for x, r in zip(sv, d)), f"{sum(1 for x in sv if x)} people")
    placed = [r for r in d if V.band(r[h["IRIS proposed level"]]) != "X"]
    wi = [h[x] for x in h if str(x).startswith("Weightage IRIS")]
    wm = [h[x] for x in h if str(x).startswith("Weightage Market")]
    check("IRIS and Market weightage for every placed person, in five steps",
          all(r[c] in codes for r in placed for c in wi + wm), f"{len(placed)} people")
    check("Market weights Experience and Education VL · no method weights them on their own",
          all(r[h["Weightage Market · Experience"]] == "VL" and r[h["Weightage Market · Education"]] == "VL" for r in placed))
    ri = [h[x] for x in h if str(x).startswith("Rating IRIS")]
    check("IRIS ratings are 1 to 5 against the level, blank where nothing was received",
          all(r[c] in (None, "", 1, 2, 3, 4, 5) for r in d for c in ri) and sum(1 for r in d for c in ri if r[c]) > 2000)
    gap = [r[h["Emp ID"]] for r in d if abs(float(r[h["Net live"]] or 0) - float(r[h["Anchor · IRIS as built"]] or 0)) > 1e-5]
    check("Before any input, the live score equals IRIS's anchor for everyone — no one moves by accident", not gap, f"{len(gap)} {gap[:3]}")
    if not SIMPLE:
        mb = list(wb[N["MB"]].iter_rows(values_only=True))
        lmh_rows = [r for r in mb if r and r[2] in codes and r[1] and r[0] != "Role group"]
        check("Market Benchmark gives five steps with the published numbers, the importance and references on every row",
              len(lmh_rows) >= 56 and all(r[3] and r[7] and r[9] and r[10] for r in lmh_rows), f"{len(lmh_rows)} rows")
    rj = list(wb[N["RJ"]].iter_rows(min_row=6, values_only=True))
    jh = {x: i for i, x in enumerate(rj[0])}
    jr = [r for r in rj[1:] if r[jh["Role ID"]]]
    per = collections.Counter(r[jh["Role ID"]] for r in jr)
    kinds = collections.Counter(r[jh["Row"]] for r in jr)
    laddered = {r[lh["Role ID"]] for r in lrows}
    lv_cols = [k for k in jh if k in V.LEVELS_PC]
    pairs = [(a, b) for a, b in zip(jr[0::2], jr[1::2])]
    cells = [(lv, b[jh[lv]]) for a, b in pairs for lv in lv_cols if a[jh[lv]] not in (None, "", "—")]
    filled = sum(1 for lv, x in cells if x not in (None, "", "—"))
    check("Role Journey: two rows per role — IRIS and Market benchmark — for every role on a ladder, market filled at its levels",
          set(per) == laddered and all(v == 2 for v in per.values()) and kinds["IRIS"] == kinds["Market benchmark"]
          and all(a[jh["Row"]] == "IRIS" and b[jh["Row"]] == "Market benchmark" for a, b in pairs) and filled == len(cells),
          f"{len(per)} roles · market at {filled} of {len(cells)} role-levels")
    ref = wb[N["REF"]]
    links = sum(1 for row in ref.iter_rows(values_only=True) for v in row if isinstance(v, str) and v.startswith("http"))
    srcs = {v.split()[0] for row in ref.iter_rows(values_only=True) for v in row[:1] if isinstance(v, str) and re.match(r"^S\d\d", v)}
    check("References & Benchmarks, last sheet: every source with its link, the crosswalk and the method sources",
          links >= 100 and {"S01", "S70", "S71"} <= srcs, f"{links} links · {len(srcs)} source IDs")
    raw = list(wb[N["RAW"]].iter_rows(min_row=6, values_only=True))
    rwh = [x for x in raw[0] if x]
    personal = [x for x in rwh if RW.fkey(x.split(" · ")[0]) in RW.DROP]
    qcols = [x for x in rwh if re.match(r"^Q\d+ · ", x)]
    check("Raw Data holds every column we received, all 20 survey questions, no personal fields",
          len(rwh) >= 150 and len(qcols) == 20 and not personal, f"{len(rwh)} columns · {len(qcols)} questions · personal {personal}")
    grp = [J.text(x) for x in list(wb[N["RAW"]].iter_rows(min_row=5, max_row=5, values_only=True))[0] if x]
    if SIMPLE:
        check("Simplified Raw Data keeps what was received; the earlier versions' analysis stays in the Master",
              not any(g.startswith(("ON FILE", "EMPLOYEE MASTER · OLDER")) for g in grp), str(grp))
    else:
        check("Raw Data also holds the analysis on file from every earlier version — placement, scoring, profile, talent, data quality",
              sum(1 for g in grp if g.startswith("ON FILE")) == 6, str([g[:30] for g in grp if g.startswith("ON FILE")]))
    status = collections.Counter(r[1] for r in raw[1:] if r[1])
    check("Raw Data holds everyone plus the people not yet mapped and those only in the source files", status.get("Mapped") == 586, dict(status))
    rh_ = {x: i for i, x in enumerate(raw[0])}
    by_id, by_nm = {}, collections.defaultdict(str)
    for r in raw[1:]:
        blob = " ".join(re.sub(r"\s+", " ", str(x).replace("\u200b", "")).lower() for x in r if x not in (None, ""))
        for x in str(r[rh_["Emp ID"]] or "").split("\n"):
            if J.key(x.split(": ")[-1]).startswith(("iris", "intn")):
                by_id.setdefault(J.key(x.split(": ")[-1]), blob)
        for x in str(r[rh_["Employee Name"]] or "").split("\n"):
            by_nm[J.nname(x.split(": ")[-1])] += " ¦ " + blob
    n_ans, lost = 0, []
    if RW.v1_files():
        for biz, hdr, r in RW.v1_rows():
            j = {x: i for i, x in enumerate(hdr)}
            b = by_id.get(J.key(r[j["Emp Id"]])) or by_nm.get(J.nname(r[j["Employee Name"]]), "")
            for q in RW.SURVEY_Q + [RW.ADDITIONAL]:
                v = RW.val(r[j[q]], q) if q in j else ""
                if v != "" and not (isinstance(v, str) and RW.PLACEHOLDER.match(v)):
                    n_ans += 1
                    if re.sub(r"\s+", " ", str(v).replace("\u200b", "")).strip().lower() not in b:
                        lost.append((r[j["Employee Name"]], q))
    check("Every survey answer received, question by question, is on the person's Raw Data row",
          n_ans > 9000 and not lost, f"{n_ans} answers · {len(lost)} missing {lost[:3]}")
    src = os.environ.get("JA_SRC", "")
    pb = {k: re.sub(r"[\W_]+", "", v) for k, v in by_id.items()}
    pn_ = {k: re.sub(r"[\W_]+", "", v) for k, v in by_nm.items()}
    n_cells, lost2 = 0, []
    for emp, name, biz, cells in RW.all_sources_rows(src):
        b = pb.get(J.key(emp)) or pn_.get(J.nname(name), "")
        for g, hh, val in cells:
            if RW.fkey(hh) in RW.DROP or (SIMPLE and g == "OLD"):   # the older Employee Master copy is in the Master file only
                continue
            v = RW.val(val, hh.split(" · ")[-1])
            if v == "" or (isinstance(v, str) and RW.PLACEHOLDER.match(v)):
                continue
            n_cells += 1
            if re.sub(r"[\W_]+", "", str(v).replace("\u200b", "")).lower() not in b:
                lost2.append((name, hh))
    check("The All-Sources collation: every non-personal cell is on the person's Raw Data row",
          n_cells > 50000 and not lost2, f"{n_cells} cells · {len(lost2)} missing {lost2[:3]}")
    with_q = sum(1 for x in sv if x and "Q1 · " in x)
    check("Inputs received in the survey on Talent Mapping carries the answers question by question", with_q >= 515, f"{with_q} people")
    pc = list(wb[N["PC"]].iter_rows(min_row=8, max_row=8, values_only=True))[0]
    check("Person Card resolves", isinstance(pc[1], str) and "·" in pc[1], str(pc[1]))
    z = zipfile.ZipFile(OUTFILE)
    xml = "".join(z.read(n).decode("utf8") for n in z.namelist() if n.startswith("xl/worksheets/sheet"))
    check("Excel-safe: no '=' in list validations, no formula errors", "<formula1>=" not in xml and 't="e"' not in xml)
    wb.close()
    glitches()
    live()
    print(f"\n{sum(RES)} of {len(RES)} checks pass")
    return 0 if all(RES) else 1


def glitches():
    """Layout and integrity: things that make a sheet look broken or make Excel repair the file."""
    wb = openpyxl.load_workbook(OUTFILE)
    names = set(wb.sheetnames)
    bad_links, long_f, tall, junk, blank_hdr, bad_dv = [], 0, [], 0, [], []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for x in row:
                v = x.value
                if isinstance(v, str):
                    if v.startswith("=") and len(v) > 8000:
                        long_f += 1
                    if v.strip() in ("None", "nan", "NaN", "#N/A") and not v.startswith("="):
                        junk += 1
                if x.hyperlink is not None and x.hyperlink.location is None and str(x.hyperlink.target or "").startswith("#"):
                    tgt = str(x.hyperlink.target)[1:].split("!")[0].strip("'")
                    if tgt not in names:
                        bad_links.append((ws.title, x.coordinate, tgt))
        for rn, rd in ws.row_dimensions.items():
            if rd.height and rd.height > 409:
                tall.append((ws.title, rn))
        if ws.title in (N["TM"], N["RL"], N["RAW"], N["RJ"]):
            hdr = [c.value for c in ws[6]]
            last = max(j for j, v in enumerate(hdr) if v) if any(hdr) else 0
            gaps = [j + 1 for j, v in enumerate(hdr[:last + 1]) if not v]
            if gaps:
                blank_hdr.append((ws.title, gaps[:5]))
        for dv in ws.data_validations.dataValidation:
            if not dv.sqref or str(dv.formula1 or "").startswith("="):
                bad_dv.append(ws.title)
    dn_bad = []
    for nm, dn in wb.defined_names.items():
        for title, coord in dn.destinations:
            if title not in names:
                dn_bad.append(nm)
    check("No glitches: links and named ranges point to real sheets, formulas within Excel's limits, row heights valid, no stray "
          "'None' text, every column titled, drop-downs valid",
          not bad_links and not long_f and not tall and not junk and not blank_hdr and not bad_dv and not dn_bad,
          f"links {bad_links[:2]} · long formulas {long_f} · tall rows {tall[:2]} · junk {junk} · blank headers {blank_hdr[:2]} · "
          f"validations {bad_dv[:2]} · names {dn_bad[:2]}")
    wb.close()


def ladder_of(lv):
    return V.LEVELS_PC if lv in ("B+", "C+") else V.LEVELS_STD


def predict(lv, weights, ratings, anchor):
    """The workbook's rule, recomputed here: shift = weighted rating − anchor; round half away from zero; at most two levels."""
    sh = SC.net(weights, ratings, PKEYS) - anchor
    st = int(abs(sh) + 0.5) * (1 if sh >= 0 else -1)
    st = max(-SC.MAX_SHIFT, min(SC.MAX_SHIFT, st))
    lad = ladder_of(lv)
    return lad[max(0, min(len(lad) - 1, lad.index(lv) + st))], sh


def live():
    work = tempfile.mkdtemp(prefix="jam_")
    base = openpyxl.load_workbook(OUTFILE, read_only=True, data_only=True)
    t0 = list(base[N["TM"]].iter_rows(min_row=6, values_only=True))
    base.close()
    h0 = {x: i for i, x in enumerate(t0[0])}
    pn = dict(PARAMS)
    rows0 = {n + 6: r for n, r in enumerate(t0[1:], 1) if r[h0["Emp ID"]]}

    def w_of(r):
        return {k: r[h0[f"Weightage in use · {pn[k]}"]] for k in PKEYS}

    def r_of(r):
        return {k: r[h0[f"Rating IRIS · {pn[k]}"]] for k in PKEYS}

    def anc(r):
        return float(r[h0["Anchor · IRIS as built"]] or 0)

    wb = openpyxl.load_workbook(OUTFILE)
    ws = wb[N["TM"]]
    hd = {c.value: c.column for c in ws[6]}
    rl = wb[N["RL"]]
    lh = {c.value: c.column for c in rl[6]}
    lrow = {rl.cell(r, lh["Key"]).value: r for r in range(7, rl.max_row + 1)}
    by_key = collections.defaultdict(list)
    for r, v in rows0.items():
        if v[h0["Ladder key"]]:
            by_key[v[h0["Ladder key"]]].append(r)
    role_of = {r: v[h0["Role key"]] for r, v in rows0.items()}
    # A · a role (ROLE row) and one of its levels (level row) in one department
    roles = collections.Counter(role_of[r] for r in rows0 if role_of[r])
    rkey = next(k for k, n in roles.most_common() if n >= 6 and len({rows0[r][h0["IRIS proposed level"]] for r in rows0 if role_of[r] == k}) >= 2)
    lvl = collections.Counter(rows0[r][h0["IRIS proposed level"]] for r in rows0 if role_of[r] == rkey).most_common(1)[0][0]
    lkey = collections.Counter(rows0[r][h0["Ladder key"]] for r in rows0 if role_of[r] == rkey).most_common(1)[0][0]
    rl.cell(lrow[rkey], lh["Leader · Complexity & Expertise"]).value = "VL"
    rl.cell(lrow[lkey], lh["Leader · Complexity & Expertise"]).value = "VH"
    # B · HR changes the IRIS base for another role and level
    bkey = next(k for k, rs in sorted(by_key.items(), key=lambda x: -len(x[1])) if len(rs) >= 3 and not k.startswith(rkey[:-4]))
    rl.cell(lrow[bkey], lh["IRIS · Leadership & Strategy"]).value = "VH"
    rl.cell(lrow[bkey], lh["IRIS · Complexity & Expertise"]).value = 15
    used = set(by_key[bkey]) | {r for r in rows0 if role_of[r] == rkey}
    # C · a leader types weightage for one person on Talent Mapping
    cand = []
    for r, v in rows0.items():
        if r in used or not v[h0["Ladder key"]] or v[h0["IRIS proposed level"]] in ("E1", "E2", "D2"):
            continue
        rs = r_of(v)
        top = max((k for k in PKEYS if rs[k] not in (None, "")), key=lambda k: rs[k], default=None)
        if top and SC.R5_MOVE[int(rs[top])] - anc(v) >= 0.5:
            cand.append((r, top))
    p_row, p_top = cand[0]
    for k in PKEYS:
        ws.cell(p_row, hd[f"Weightage Leader · {pn[k]}"]).value = "VH" if k == p_top else 0
    used.add(p_row)
    # D · a leader rates every parameter 5 for a B1, and sets an override for an A3
    t1 = next(r for r, v in rows0.items() if v[h0["IRIS proposed level"]] == "B1" and r not in used and v[h0["Ladder key"]])
    for k in PKEYS:
        ws.cell(t1, hd[f"Rating Leader · {pn[k]}"]).value = 5
    used.add(t1)
    t2 = next(r for r, v in rows0.items() if v[h0["IRIS proposed level"]] == "A3" and r not in used)
    ws.cell(t2, hd["Leader override · any level"]).value = "E2"
    used.add(t2)
    # F · a leader overrides a Band C manager down to B3: their Band B reportees are held in Band A by the band rule
    inf0 = {r: rows0[r][h0["Level in force · after the band rule"]] for r in rows0}
    emp_row = {rows0[r][h0["Emp ID"]]: r for r in rows0}
    kids = collections.defaultdict(list)
    for r in rows0:
        m = emp_row.get(rows0[r][h0["Manager Emp ID"]])
        if m:
            kids[m].append(r)
    mgr_c = next(r for r in rows0 if r not in used and rows0[r][h0["IRIS proposed level"]] in ("C1", "C2")
                 and sum(1 for k in kids[r] if V.band(inf0[k]) == "B") >= 2 and not any(k in used for k in kids[r]))
    ws.cell(mgr_c, hd["Leader override · any level"]).value = "B3"
    used.add(mgr_c)
    # E · a vacant level
    vac = next(r for r in range(7, rl.max_row + 1) if rl.cell(r, lh["People in IRIS"]).value == "Vacant" and rl.cell(r, lh["Row"]).value == "level")
    rl.cell(vac, lh["Leader · Business Impact"]).value = "H"
    wb.save(os.path.join(work, "e.xlsx"))
    out = X.recalc_copy(os.path.join(work, "e.xlsx"), work)
    v = openpyxl.load_workbook(out, read_only=True, data_only=True)
    t = list(v[N["TM"]].iter_rows(min_row=6, values_only=True))
    h = {x: i for i, x in enumerate(t[0])}
    got = {r: t[r - 6] for r in rows0}

    def expect(r, wts):
        return predict(rows0[r][h0["IRIS proposed level"]], wts, r_of(rows0[r]), anc(rows0[r]))[0]

    # D
    w1 = w_of(rows0[t1])
    e1 = predict("B1", w1, {k: 5 for k in PKEYS}, anc(rows0[t1]))[0]
    check("Live: a leader rating every parameter 5 moves a B1 up — to exactly the level the rule gives",
          got[t1][h["Leader proposed level"]] == e1 and e1 != "B1", f"{got[t1][h['Leader proposed level']]} · expected {e1}")
    check("Live: an override to E2 shows as the leader's proposal, and the band rule holds the level in force below the manager's band",
          got[t2][h["Leader proposed level"]] == "E2" and got[t2][h["Level in force · after the band rule"]] != "E2"
          and str(got[t2][h["Band rule · live"]]).startswith("Held at") and not str(got[t2][h["Manager check · live"]]).startswith("⚠"),
          f"{got[t2][h['Level in force · after the band rule']]} · {got[t2][h['Band rule · live']]} · {got[t2][h['Manager check · live']]}")
    # C
    wC = {k: ("VH" if k == p_top else 0) for k in PKEYS}
    eC = expect(p_row, wC)
    check("Live: weightage typed for one person on Talent Mapping re-scores them and changes their level and band",
          got[p_row][h["Leader proposed level"]] == eC and eC != rows0[p_row][h0["IRIS proposed level"]]
          and got[p_row][h["Leader proposed band"]] not in (None, ""),
          f"{rows0[p_row][h0['IRIS proposed level']]} → {got[p_row][h['Leader proposed level']]} · expected {eC} · band {got[p_row][h['Leader proposed band']]}")
    # A
    in_role = [r for r in rows0 if role_of[r] == rkey]
    bad_a, moved_a = [], 0
    for r in in_role:
        wts = w_of(rows0[r])
        wts["C"] = "VH" if rows0[r][h0["Ladder key"]] == lkey else "VL"
        e = expect(r, wts)
        moved_a += e != rows0[r][h0["IRIS proposed level"]]
        if got[r][h["Leader proposed level"]] != e or got[r][h["Weightage in use · Complexity & Expertise"]] != wts["C"]:
            bad_a.append((rows0[r][h0["Employee name"]], got[r][h["Leader proposed level"]], e))
    check("Live: a ROLE-row weightage reaches everyone in the department's role, a level row wins for its level, and each "
          "person lands exactly where the rule puts them", not bad_a and in_role, f"{len(in_role)} people · {moved_a} change level · {bad_a[:2]}")
    # B
    bad_b, moved_b = [], 0
    for r in by_key[bkey]:
        wts = w_of(rows0[r])
        wts["L"], wts["C"] = "VH", 15
        e = expect(r, wts)
        moved_b += e != rows0[r][h0["IRIS proposed level"]]
        if got[r][h["Leader proposed level"]] != e:
            bad_b.append((rows0[r][h0["Employee name"]], got[r][h["Leader proposed level"]], e))
    check("Live: HR changing the IRIS base on the Role Ladder (a code and a number) re-scores everyone at that role and level",
          not bad_b, f"{len(by_key[bkey])} people · {moved_b} change level · {bad_b[:2]}")
    held = [k for k in kids[mgr_c] if V.band(inf0[k]) == "B"]
    ok_f = all(got[k][h["Level in force · after the band rule"]] == "A3" and str(got[k][h["Band rule · live"]]).startswith("Held at A3")
               for k in held) and got[mgr_c][h["Level in force · after the band rule"]] == "B3"
    check("Live: the band rule — a manager moved down to B3 holds their Band B reportees at A3, flagged; the only warnings left are "
          "where no band is left below an A-level manager",
          ok_f and all(str(got[r][h["Manager check · live"]]).startswith("⚠ No band left") for r in rows0
                       if str(got[r][h["Manager check · live"]]).startswith("⚠")),
          f"{len(held)} reportees held · {[got[k][h['Level in force · after the band rule']] for k in held][:5]} · "
          f"{[got[k][h['Band rule · live']] for k in held][:1]} · manager {got[mgr_c][h['Level in force · after the band rule']]} · "
          f"warnings {[(got[r][h['Employee name']], got[r][h['Manager check · live']]) for r in rows0 if str(got[r][h['Manager check · live']]).startswith('⚠')][:3]}")
    others = sum(1 for r in rows0 if got[r][h["Leader proposed level"]] not in (None, "") and r not in used and r not in by_key[bkey])
    check("Live: nobody else gets a leader proposal", others == 0, f"{others}")
    lr = list(v[N["RL"]].iter_rows(min_row=6, values_only=True))
    lx = {x: i for i, x in enumerate(lr[0])}
    vr = lr[vac - 6]
    check("Live: a vacant level can be weighted too · In use and the shares follow",
          vr[lx["In use · Business Impact"]] == "H" and vr[lx["Changes here"]] == 1
          and abs(sum(float(vr[lx[f"Share · {pn[k]}"]] or 0) for k in PKEYS) - 1) < 1e-6)
    br = lr[lrow[bkey] - 6]
    check("Live: the Role Ladder counts HR's base changes", br[lx["Changes here"]] == 2, str(br[lx["Changes here"]]))
    errs = sum(1 for ws_ in v.worksheets for row in ws_.iter_rows(values_only=True) for x in row if isinstance(x, str) and re.match(r"^#(REF|N/A|VALUE|NAME|DIV)", x))
    check("Live: no formula errors after the edits", errs == 0, str(errs))
    v.close()
    shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
