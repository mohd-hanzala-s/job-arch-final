"""HR Super Master — Final v1.13.

Starts from the v1.12.2 HR Super Master and keeps every sheet, value, formula,
chart, table, filter and named range. Adds:
  · sheet 2 — KRAs received from managers; what the employee shared in the survey;
              links and data quality (Role ID, manager's Emp ID and level, data
              quality, weight sensitivity)
  · sheet 3 — values the v1.12.2 clean-up lost, restored from the source files,
              and the CARBON Mapping Master columns it never carried
  · 1b Person Card — the same three sections
  · 20. Not Yet Mapped — people in the source files who are not in the system
  · Start and Contents — what v1.13 adds

Usage: python3 build_super.py <folder with the five talent-mapping source files>
"""
import copy
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import TableColumn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ja_common as J  # noqa: E402
import ja_person as P  # noqa: E402
import ja_xlsx as X  # noqa: E402

OUTFILE = os.path.join(J.OUT, "IRIS_JA_2_HR_Super_Master_v1.13.xlsx")

NAVY = "1F3864"
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
RED_FONT = Font(name="Calibri", size=9, color="9C0006")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
MED = Side(style="medium", color="1F3864")
GROUPS = {
    "kra": ("KRAs RECEIVED FROM MANAGERS — as submitted · source talent-mapping files, October 2026", "375623", "E2EFDA"),
    "survey": ("WHAT THE EMPLOYEE SHARED IN THE SURVEY — per person · CARBON Mapping Master; other products need the How We Grow export", "7030A0", "EFE6F7"),
    "links": ("LINKS & DATA QUALITY — keys to the role profile and manager, data checks, weight sensitivity", "595959", "EDEDED"),
}
LINK_COLS = [
    ("Role ID · catalogue", 12),
    ("Reporting manager · Emp ID", 14),
    ("Reporting manager · final level", 11),
    ("Emp ID in the source files · where different", 16),
    ("Data quality", 42),
    ("Weight sensitivity", 30),
]


def is_needed(v):
    s = J.text(v)
    return s.startswith(("KRAs not received", "Full survey answers not in", "No survey response"))


def person_values(M, F, p):
    k = P.kras(M, p)
    s8 = F.s8row(p)
    sv = P.survey(M, p, s8)
    cat = F.cat_of[p["i"]]
    mgr = F.manager(p)
    links = [cat[0] if cat else "", mgr[0], mgr[1], ", ".join(sorted(p["src_ids"])),
             F.quality(p, k[5], sv[-1]), F.sensitivity(p)]
    return k, sv, links


def style_header(c, fill_hex, size=9):
    c.font = Font(name="Calibri", size=size, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=fill_hex)
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
    c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_body(c, tint_hex, first=False):
    c.font = Font(name="Calibri", size=9)
    c.fill = PatternFill("solid", fgColor=tint_hex)
    c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
    c.border = Border(left=MED if first else THIN, right=THIN, top=THIN, bottom=THIN)


def add_columns(ws, table_name, header_row, first_row, cols_by_group, rows_values):
    """Append column groups at the right of a sheet's table. rows_values[i] is the list of values for data row i."""
    start = ws.max_column + 1
    c = start
    spec = []
    for gkey, cols in cols_by_group:
        title, dark, tint = GROUPS[gkey]
        g = ws.cell(header_row - 1, c, title)
        style_header(g, dark, 10)
        for j, (h, w) in enumerate(cols):
            ws.column_dimensions[L(c)].width = w
            hc = ws.cell(header_row, c, h)
            style_header(hc, dark)
            if j:
                ws.cell(header_row - 1, c).fill = PatternFill("solid", fgColor=dark)
            spec.append((c, tint, j == 0))
            c += 1
    for i, vals in enumerate(rows_values):
        r = first_row + i
        for (col, tint, first), v in zip(spec, vals):
            cell = ws.cell(r, col, v if v != "" else None)
            style_body(cell, tint, first)
            if is_needed(v):
                cell.fill, cell.font = RED_FILL, RED_FONT
    # extend the Excel table and its filter
    t = ws.tables[table_name]
    a, b = t.ref.split(":")
    last_row = "".join(ch for ch in b if ch.isdigit())
    t.ref = f"{a}:{L(c - 1)}{last_row}"
    if t.autoFilter is not None:
        t.autoFilter.ref = t.ref
    nid = max(tc.id for tc in t.tableColumns) + 1
    for col in range(start, c):
        t.tableColumns.append(TableColumn(id=nid, name=ws.cell(header_row, col).value))
        nid += 1
    return start, c - 1


def stamp(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=3):
            for cell in row:
                if isinstance(cell.value, str) and "Final v1.12 · October 2026" in cell.value:
                    cell.value = cell.value.replace("Final v1.12 · October 2026", "Final v1.13 · October 2026")


def person_card(wb, start_col):
    ws = wb["1b. Person Card"]
    hdr_src, lab_src, val_src = ws["A53"], ws["A54"], ws["B54"]
    r = ws.max_row + 2
    tm = "'2. Talent Mapping'"
    col = start_col

    def section(title):
        nonlocal r
        c = ws.cell(r, 1, title)
        c._style = copy.copy(hdr_src._style)
        for k in range(2, 9):
            ws.cell(r, k)._style = copy.copy(hdr_src._style)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        r += 1

    def line(label, height):
        nonlocal r, col
        a = ws.cell(r, 1, label)
        a._style = copy.copy(lab_src._style)
        L_ = L(col)
        f = (f'=IFERROR(IF(INDEX({tm}!${L_}$7:${L_}$592,$J$7)="","—",'
             f'INDEX({tm}!${L_}$7:${L_}$592,$J$7)),"")')
        b = ws.cell(r, 2, f)
        b._style = copy.copy(val_src._style)
        for k in range(3, 9):
            ws.cell(r, k)._style = copy.copy(val_src._style)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        ws.row_dimensions[r].height = height
        r += 1
        col += 1

    section("KRAs RECEIVED FROM MANAGER")
    for h, _ in P.KRA_COLS:
        line(h, 170 if "KRAs" in h and "status" not in h else (60 if "JD" in h else 30))
    section("WHAT THE EMPLOYEE SHARED IN THE SURVEY")
    for h, _ in P.SURVEY_COLS:
        line(h, 30 if "answered" in h or "status" in h else 120)
    section("LINKS & DATA QUALITY")
    for h, _ in LINK_COLS:
        line(h, 40 if h == "Data quality" else 30)
    ws.print_area = f"A1:H{r - 1}"
    return r - 1


def not_yet_mapped(wb, M, S, people, unmapped):
    idx = wb.sheetnames.index("19. Competency Library") + 1
    ws = wb.create_sheet("20. Not Yet Mapped", idx)
    tmpl = wb["17. Data Gaps"]
    for rr in (1, 2, 3):
        ws.cell(rr, 1)._style = copy.copy(tmpl.cell(rr, 1)._style)
    ws["A1"] = "Not yet mapped — people in the source files who are not in the job architecture"
    ws["A2"] = ("Sheet 20 · Found in the managers' talent-mapping files of October 2026 but not among the 586 people placed in "
                "v1.12.2. Placements change only on explicit request, so they are listed, not placed. HR: decide for each — add "
                "and place, confirm they have left, or confirm they are out of scope.")
    ws["A3"] = J.STAMP
    ws["H1"] = "◂ Contents"
    ws["H1"].hyperlink = "#'Contents'!A1"
    ws["H1"].font = Font(color="0563C1", underline="single", size=9)
    heads = [("Source file", 26), ("Emp ID", 12), ("Employee name", 24), ("Business · per source", 14), ("Department", 18),
             ("Designation", 24), ("Reporting manager", 20), ("Category · status", 16), ("Date of joining", 12),
             ("Total experience", 14), ("KRAs set by manager · annual", 50), ("KRAs · quarterly", 40),
             ("KRAs · self-reported annual", 40), ("Rating · potential · 9-box", 22),
             ("What they shared in the survey · CARBON Mapping Master", 60), ("Actual JD · from manager", 35),
             ("Possible match in this file", 26), ("Status / HR action", 40)]
    for j, (h, w) in enumerate(heads, 1):
        style_header(ws.cell(5, j, h), NAVY)
        ws.column_dimensions[L(j)].width = w
    # merge records of the same person across files
    merged = {}
    for u in unmapped:
        rec, s = u["rec"], u["source"]
        k = J.key(rec["id"]) or J.nname(rec["name"])
        merged.setdefault(k, []).append((s, rec))
    names = [(J.nname(p["name"]), p) for p in people]
    r = 6
    for k, recs in merged.items():
        def g(*hs):
            for s, rec in recs:
                for h, v in rec["row"].items():
                    if J.hkey(h) in [J.hkey(x) for x in hs] and J.has(v):
                        return J.clean(v)
            return ""
        nm = recs[0][1]["name"]
        biz = {"MM": "CARBON", "CRAW": "CARBON"}.get(recs[0][0]["kind"], g("Business Line") or recs[0][0]["label"].split("· ")[-1])
        survey = "\n\n".join(x for x in (g("Role Mandate"), g("Reported Deliverables & Responsibilities"), g("Current Capabilities"),
                                         g("Development Priorities & Learning Needs"), g("Challenges & Support Required")) if x)
        t = J.nname(nm).split()
        match = [p for n, p in names if n.split()[:1] == t[:1] and (n.split()[-1:] == t[-1:] or len(t) == 2 and t[1] in n.split())]
        removed = J.nname(nm) == "SHREYASH PAWAR"
        status = ("Removed from the system by HR · review of 9 October 2026" if removed else
                  "Not in the v1.12.2 people list · HR to decide: add and place, confirm left, or out of scope")
        vals = [" + ".join(sorted({f"{s['file']} · {s['sheet']}" for s, _ in recs})), recs[0][1]["id"] or "—", nm, biz,
                g("Department"), g("Designation", "Current Designation", "New Designation"),
                g("Reporting Manager", "Current Reporting Manager"), " · ".join(x for x in (g("Employee Category", "Category"), g("Employee Status")) if x),
                g("DoJ"), g("Total Experience"),
                g("Annual KRAs (objective Measurable) Minimum 5 - set by Manager", "KRAs Submitted"), g("Quarterly KRAs"),
                g("Annual KRAs (objective Measurable) Minimum 5 - Self"),
                " · ".join(x for x in (g("Quarterly Performance Rating (1-5)"), g("Potential Rating (High/Med/Low)"), g("9-Box Placement")) if x),
                survey, g("Actual JD"),
                "; ".join(f"{p['name']} · {p['emp']} · {p['biz']}" for p in match[:3]), status]
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v or None)
            style_body(c, "FFFFFF", j == 1)
        ws.row_dimensions[r].height = 90
        r += 1
    ws.freeze_panes = "D6"
    ws.auto_filter.ref = f"A5:{L(len(heads))}{r - 1}"
    return len(merged)


def contents_and_start(wb, n_unmapped, stats):
    ws = wb["Contents"]
    last = ws.max_row
    src = [ws.cell(last, j) for j in (1, 2, 3)]
    r = last + 1
    vals = [int(J.text(ws.cell(last, 1).value) or 0) + 1, "20. Not Yet Mapped",
            f"{n_unmapped} people found in the managers' talent-mapping files of October 2026 who are not in the job architecture: "
            "their data, KRAs and survey summary, a possible match, and the HR action."]
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v)
        c._style = copy.copy(src[j - 1]._style)
    ws.cell(r, 2).hyperlink = "#'20. Not Yet Mapped'!A1"
    ws.row_dimensions[r].height = ws.row_dimensions[last].height
    # v1.13 note on the Contents and the Start sheet
    st = wb["Start — How It Works"]
    r = st.max_row + 2
    h5, h6, body = st["A5"], st["A6"], st["A7"]
    c = st.cell(r, 1, "What v1.13 adds — 12 October 2026")
    c._style = copy.copy(h5._style)
    r += 1
    for j, h in enumerate(["#", "Added", "What it holds", "Where"], 1):
        st.cell(r, j, h)._style = copy.copy(h6._style)
    r += 1
    rows = [
        ("KRAs received from managers", f"Manager-set annual KRAs, quarterly KRAs, self-reported KRAs, the Q1 tentative rating and the actual JD, as submitted. On file for {stats['kra']} of 586 people; the rest are red with 'KRAs not received · manager to supply'. Text the earlier clean-up had cut from inside brackets is restored, written with ' · ' instead of brackets.", "Sheet 2 · 1b Person Card"),
        ("What the employee shared in the survey", f"Per person: role mandate and core purpose, deliverables and accountability, capabilities and tools, development and learning needs, challenges and support — {stats['survey']} CARBON people from the CARBON Mapping Master. For the other products the full answers are not in the files received: the cell says so in red and the How We Grow export is needed. Also on every row: questions answered, the survey and KRA evidence used for each rating, the role-holders' themes and the learning topics named.", "Sheet 2 · 1b Person Card"),
        ("Links and data quality", "Role ID to the role catalogue, reporting manager's Emp ID and final level, Emp ID where the source files use a different one, a data-quality line and weight sensitivity for every person.", "Sheet 2 · 1b Person Card"),
        ("Raw data completed", f"{stats['restored']} people have values restored from the source files where v1.12.2 had a blank or a shortened value · listed per person. The CARBON Mapping Master's comments and five survey-summary columns are added.", "Sheet 3"),
        ("Not yet mapped", f"{n_unmapped} people in the source files who are not in the job architecture, for HR to decide.", "Sheet 20"),
        ("Privacy rule changed by HR", "Per-person survey content is now shown against names in this file, by HR's decision of 12 October 2026. The file is for HR only. Personal fields — gender, date of birth, age, e-mail, resignation and relieving dates — are still never written, and pay stays blank.", "Whole file"),
    ]
    for i, row in enumerate(rows, 1):
        for j, v in enumerate((i,) + row, 1):
            c = st.cell(r, j, v)
            c._style = copy.copy(body._style)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        st.row_dimensions[r].height = 75
        r += 1


def main(srcdir):
    os.makedirs(J.OUT, exist_ok=True)
    M, S, people, unmapped = J.build(srcdir)
    F = P.Folds(M, people)
    wb = openpyxl.load_workbook(J.MASTER)

    # sheet 2
    ws2 = wb["2. Talent Mapping"]
    rows = []
    stats = {"kra": 0, "survey": 0, "restored": 0}
    for p in people:
        k, sv, links = person_values(M, F, p)
        stats["kra"] += k[5].startswith("Received")
        stats["survey"] += sv[-1].startswith("Per-person")
        rows.append(k + sv + links)
    assert [J.text(ws2.cell(7 + i, 2).value) for i in range(len(people))] == [p["emp"] for p in people]
    start, end = add_columns(ws2, "T_2__Talent_Mapping", 6, 7,
                             [("kra", P.KRA_COLS), ("survey", P.SURVEY_COLS), ("links", LINK_COLS)], rows)
    for rr in range(7, 7 + len(people)):
        ws2.row_dimensions[rr].height = max(ws2.row_dimensions[rr].height or 0, 340)

    # sheet 3
    ws3 = wb["3. Raw Data & Market Standards"]
    raw_rows = []
    for i, p in enumerate(people):
        row, extra, restored = P.restore_raw(M, p)
        if restored:
            stats["restored"] += 1
            for j, h in enumerate(M["raw_hdr"]):
                if h and any(h.endswith(x) for x in restored) and row[j] != p["raw"][j]:
                    ws3.cell(7 + i, j + 1).value = row[j]
        raw_rows.append(extra)
    GROUPS["raw"] = ("ADDED IN v1.13 — CARBON Mapping Master columns v1.12.2 did not carry · values restored from the source files", "548235", "EEF6EA")
    add_columns(ws3, "T_3__Raw_Data___Market_Standards", 6, 7,
                [("raw", [(h, 45) for _, _, h in P.EXTRA_RAW] + [(P.RAW_LOG, 40), (P.RAW_IDS, 16)])], raw_rows)

    person_card(wb, start)
    n_un = not_yet_mapped(wb, M, S, people, unmapped)
    contents_and_start(wb, n_un, stats)
    stamp(wb)
    wb.save(OUTFILE)
    work = os.path.join(os.environ.get("JA_WORK", "/tmp"), "ja_super")
    os.makedirs(work, exist_ok=True)
    n_f, errors = X.cache_values(OUTFILE, work)
    print(f"Wrote {OUTFILE}: {len(people)} people · KRAs {stats['kra']} · CARBON survey {stats['survey']} · "
          f"raw restored {stats['restored']} · not yet mapped {n_un} · formulas {n_f} · formula errors {len(errors)}")
    if errors:
        print(errors[:20])
    return OUTFILE


if __name__ == "__main__":
    main(sys.argv[1])
