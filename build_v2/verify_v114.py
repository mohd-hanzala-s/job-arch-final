"""Checks for the three v1.14 workbooks — HR Super Master, HR Master Simplified, Business-Head Workbook.
Exit code 0 = every check passes. Usage: python3 verify_v114.py
"""
import collections
import os
import re
import sys
import zipfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ja_common as J  # noqa: E402
import ja_v114 as V  # noqa: E402
import ja_xlsx as X  # noqa: E402
import shutil, subprocess, tempfile

SUPER = os.path.join(J.OUT, f"IRIS_JA_2_HR_Super_Master_{J.VER}.xlsx")
SIMPL = os.path.join(J.OUT, f"IRIS_JA_2S_HR_Master_Simplified_{J.VER}.xlsx")
BH = os.path.join(J.OUT, f"IRIS_JA_7_Business_Head_Workbook_{J.VER}.xlsx")
RESULTS = []


def check(name, ok, detail=""):
    RESULTS.append(ok)
    print(("PASS " if ok else "FAIL ") + name + (f" — {detail}" if detail else ""))


def num_eq(a, b):
    return isinstance(a, (int, float)) and isinstance(b, (int, float)) and abs(a - b) < 1e-6


def titled(v, w):
    if not isinstance(v, str) or not isinstance(w, str):
        return False
    x = V.TITLE_CELL.get(v.strip(), v)
    for o, r in V.TITLE_REPLACE:
        x = x.replace(o, r)
    return x == w or V.TITLE_EXACT.get(v.strip()) == w


def main():
    orig = openpyxl.load_workbook(J.MASTER, read_only=True, data_only=True)
    sup = openpyxl.load_workbook(SUPER, read_only=True, data_only=True)
    sim = openpyxl.load_workbook(SIMPL, read_only=True, data_only=True)

    # 1 every original sheet kept, every original value kept (sheet 3 restorations and stamps aside)
    missing = [n for n in orig.sheetnames if n not in sup.sheetnames]
    check("Super Master keeps every v1.12.2 sheet", not missing, f"{len(orig.sheetnames)} sheets · missing {missing}")
    changed = collections.Counter()
    decided = collections.Counter()
    for n in orig.sheetnames:
        a = list(orig[n].iter_rows(values_only=True))
        b = list(sup[n].iter_rows(values_only=True))
        for i, ra in enumerate(a):
            rb = b[i] if i < len(b) else ()
            for j, v in enumerate(ra):
                w = rb[j] if j < len(rb) else None
                if v != w and not num_eq(v, w):
                    if i < 3 and isinstance(v, str) and "Final v1.12" in v:
                        continue
                    if titled(v, w):
                        decided["titles"] += 1
                        continue
                    if w in (V.HEAD_FAMILY, V.HEAD_SUB) and n in ("1a. Placements", "1b. Person Card", "2. Talent Mapping", "3. Raw Data & Market Standards"):
                        decided["heads"] += 1
                        continue
                    changed[n] += 1
    allowed = {"3. Raw Data & Market Standards"}
    bad = {k: v for k, v in changed.items() if k not in allowed}
    check("Super Master: every v1.12.2 value unchanged, except restored raw values and the v1.14 title and family decisions", not bad,
          f"restored on sheet 3: {changed.get('3. Raw Data & Market Standards', 0)} · titles {decided['titles']} · heads' family cells {decided['heads']} · other {bad}")

    # 2 new columns present and filled
    tm = list(sup["2. Talent Mapping"].iter_rows(min_row=6, values_only=True))
    h = {x: i for i, x in enumerate(tm[0])}
    need = ["KRAs set by manager · annual", "Survey questions answered · of 20", "Role mandate & core purpose · as shared",
            "Survey status", "Role ID · catalogue", "Data quality", "Weight sensitivity"]
    check("Super Master sheet 2 carries the KRA, survey and link columns", all(x in h for x in need), f"{len(tm[0])} columns")
    rows = [r for r in tm[1:] if r[1] or r[2]]
    kra = sum(1 for r in rows if not str(r[h["KRAs set by manager · annual"]] or "").startswith("KRAs not received"))
    srv = sum(1 for r in rows if str(r[h["Survey status"]] or "").startswith("Per-person"))
    check("KRAs received shown per person; missing ones marked red with who supplies them", kra > 400, f"{kra} of {len(rows)} on file")
    check("Survey summary per person where the source holds one", srv == 205, f"{srv} CARBON people")
    check("Every person has a Role ID", all(r[h["Role ID · catalogue"]] for r in rows), f"{len(rows)} people")

    # 3 privacy and pay
    for label, wb in (("Super Master", sup), ("Simplified", sim)):
        bad_cols = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(min_row=1, max_row=8, values_only=True):
                for v in row:
                    if isinstance(v, str) and J.hkey(v) in J.PERSONAL and len(v) < 30:
                        bad_cols.append((ws.title, v))
        check(f"{label}: no personal-field columns · gender, DoB, age, e-mail, resignation, relieving", not bad_cols, str(bad_cols[:5]))
    t = list(sim["2. Talent Mapping"].iter_rows(min_row=6, values_only=True))
    hs = {x: i for i, x in enumerate(t[0])}
    pay = [x for x in t[0] if x and ("pay" in str(x).lower() and "CTC" in str(x) or str(x).startswith("Proposed range") or x == "Compa-ratio")]
    filled = sum(1 for r in t[1:] for x in pay if r[hs[x]] not in (None, ""))
    check("Simplified: pay columns blank", filled == 0, f"{len(pay)} columns, {filled} filled")

    # 4 simplified = super master, person by person
    srows = [r for r in t[1:] if r[hs["Emp ID"]] or r[hs["Employee name"]]]
    check("Simplified sheet 2 has every person, in the same order", [r[hs["Emp ID"]] for r in srows] == [r[1] for r in rows], f"{len(srows)} people")
    mism = 0
    scored = 0
    for a, b in zip(rows, srows):
        sa, sb = a[h["WEIGHTED SCORE = Σ points"]], b[hs["WEIGHTED SCORE = Σ points"]]
        if isinstance(sa, (int, float)):
            scored += 1
            if not (isinstance(sb, (int, float)) and abs(sa - sb) < 0.005):
                mism += 1
        for k in ("Level the score indicates", "FINAL LEVEL", "Final designation"):
            if a[h[k]] != b[hs[k]]:
                mism += 1
    check("Simplified: live scores, indicated level, final level and designation equal the Super Master", mism == 0, f"{scored} scored · {mism} differ")
    lost = sum(1 for k, i in h.items() if k in hs for a, b in zip(rows, srows) if a[i] not in (None, "") and b[hs[k]] in (None, ""))
    absent = [k for k in h if k not in hs]
    check("Simplified: every Super Master sheet-2 column and value is on sheet 2", not absent and lost == 0, f"missing columns {absent} · lost values {lost}")
    A8 = list(sup["8. Talent, Learning & Mobility"].iter_rows(min_row=6, values_only=True))
    lost8 = 0
    idx = {(str(r[hs["Emp ID"]]), str(r[hs["Employee name"]])): r for r in srows}
    for r in A8[1:]:
        if r[1] or r[2]:
            q = idx[(str(r[1]), str(r[2]))]
            lost8 += sum(1 for j, k in enumerate(A8[0]) if k in hs and r[j] not in (None, "") and q[hs[k]] in (None, ""))
    check("Simplified: every sheet-8 talent, competency, learning and mobility value is on sheet 2", lost8 == 0, f"lost {lost8}")
    RA = list(sup["3. Raw Data & Market Standards"].iter_rows(min_row=6, values_only=True))
    RB = list(sim["3. Raw Data"].iter_rows(min_row=6, values_only=True))
    lostr = sum(1 for a, b in zip(RA[1:], RB[1:]) for j, v in enumerate(a) if v not in (None, "") and (j + 1 >= len(b) or b[j + 1] in (None, "")))
    status = collections.Counter(r[0] for r in RB[1:] if r[0])
    check("Simplified: every raw value is on sheet 3, plus the people not yet mapped", lostr == 0, f"lost {lostr} · {dict(status)}")
    rp = [r for r in sim["5. Role Profiles"].iter_rows(min_row=7, values_only=True) if r[0]]
    hp = {x: i for i, x in enumerate(next(sim["5. Role Profiles"].iter_rows(min_row=6, max_row=6, values_only=True)))}
    holders = sum(len(str(r[hp["Holders · name · Emp ID · business"]] or "").split("\n")) for r in rp if r[hp["Holders · name · Emp ID · business"]])
    check("Role Profiles: every catalogue role plus the v1.14 Executive Leadership · Business Head roles, every holder", len(rp) >= 175 and holders == len(rows), f"{len(rp)} roles · {holders} holders")
    pc = list(sim["4. Person Card"].iter_rows(min_row=8, max_row=8, values_only=True))[0]
    check("Person Card resolves the default person", isinstance(pc[1], str) and "IRIS" in pc[1], str(pc[1]))

    # 5 workbook hygiene
    for label, path in (("Super Master", SUPER), ("Simplified", SIMPL)):
        z = zipfile.ZipFile(path)
        dv_eq = 0
        f_err = 0
        brackets = 0
        for n in z.namelist():
            if n.startswith("xl/worksheets/sheet") and n.endswith(".xml"):
                x = z.read(n).decode("utf8")
                dv_eq += len(re.findall(r"<formula1>=", x))
                f_err += len(re.findall(r't="e"', x))
        check(f"{label}: no '=' at the start of a list validation · Excel repair rule", dv_eq == 0, f"{dv_eq}")
        check(f"{label}: no formula returns an error", f_err == 0, f"{f_err}")
        charts = [n for n in z.namelist() if n.startswith("xl/charts/chart")]
        check(f"{label}: charts present", len(charts) >= (8 if label == "Super Master" else 1), f"{len(charts)}")
    new_text = 0
    for ws in sim.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str) and not v.startswith("=") and re.search(r"[()]", v):
                    new_text += 1
    check("Simplified: no text in round brackets", new_text == 0, f"{new_text} cells")
    v114_checks(sup, sim)
    print(f"\n{sum(RESULTS)} of {len(RESULTS)} checks pass")
    return 0 if all(RESULTS) else 1




def v114_checks(sup, sim):
    tm = list(sup["2. Talent Mapping"].iter_rows(min_row=6, values_only=True))
    h = {x: i for i, x in enumerate(tm[0])}
    rows = [r for r in tm[1:] if r[1] or r[2]]
    heads = [r for r in rows if r[h["Job family"]] == V.HEAD_FAMILY and r[h["Sub-family · what they do"]] == V.HEAD_SUB]
    moved = [r for r in rows if J.text(r[h["Job family & sub-family · v1.14 change"]]).startswith("Department head")]
    check("Department and function heads sit in Executive Leadership · Business Head", len(moved) == 18 and all(r in heads for r in moved),
          f"{len(moved)} heads moved · {len(heads)} people now in Executive Leadership · Business Head")
    hm = next(r for r in rows if r[h["Employee name"]] == "Hemant Karale")
    check("HR's edit kept: Hemant Karale's sub-family is Business Head", hm[h["Sub-family · what they do"]] == V.HEAD_SUB)
    lc = sup["1c. Scoring Made Simple"]
    titles = {J.text(r[1]): J.text(r[3]) for r in lc.iter_rows(values_only=True) if r[1] in ("D2", "D1", "B1")}
    want = {"D2": "Senior Vice President / Business Head", "D1": "Associate Vice President / Distinguished Principal",
            "B1": "Team Lead / Specialist / Asst Manager"}
    check("HR's title edits on the ladder: D2, D1 and B1", all(titles.get(k) == v for k, v in want.items()), str(titles))
    stale = 0
    for ws in sup.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str) and ("Associate Vice President / Business Head" in v or "Chief Specialist / Distinguished Principal" in v):
                    stale += 1
    check("No stale D1 / D2 titles anywhere in the Super Master", stale == 0, f"{stale} cells")
    # manager rule on the HR proposed level
    idx = {r[h["Emp ID"]]: r for r in rows}
    lvl = h["HR proposed level · v1.14"]
    bad = 0
    pairs = 0
    for r in rows:
        m = idx.get(r[h["Reporting manager · Emp ID"]])
        if m is None:
            continue
        a, b = J.text(r[lvl]), J.text(m[lvl])
        if "X" in (V.band(a), V.band(b)) or not a or not b:
            continue
        pairs += 1
        if V.band(a) == "E" and V.band(b) == "E":
            continue
        if V.VAL[b] >= 7 and V.BAND_ORDER.index(V.band(a)) >= V.BAND_ORDER.index(V.band(b)):
            bad += 1
        elif V.VAL[a] >= V.VAL[b]:
            bad += 1
    check("Top-down manager rule holds on every HR proposed level · heads a band apart, others a level apart", bad == 0, f"{pairs} pairs · {bad} breaks")
    st = list(sim["2. Talent Mapping"].iter_rows(min_row=6, values_only=True))
    sh = {x: i for i, x in enumerate(st[0])}
    check("Simplified carries the HR proposed level for everyone", "HR proposed level · v1.14" in sh and
          all(a[lvl] == b[sh["HR proposed level · v1.14"]] for a, b in zip(rows, [r for r in st[1:] if r[sh["Emp ID"]] or r[sh["Employee name"]]])))
    # business-head workbook
    bh = openpyxl.load_workbook(BH, read_only=True, data_only=True)
    t = list(bh["2. Talent Mapping"].iter_rows(min_row=6, values_only=True))
    th = {x: i for i, x in enumerate(t[0])}
    tr = [r for r in t[1:] if r[th["Emp ID"]] or r[th["Employee name"]]]
    hr = th["HR proposed band · market benchmark and people inputs"]
    check("Business-head workbook: every person, top down from the CEO", len(tr) == 586 and tr[0][th["Employee name"]] == "Balachandran K", f"{len(tr)} people")
    supl = {(r[h["Emp ID"]], r[h["Employee name"]]): r[lvl] for r in rows}
    check("Business-head HR proposed band = Super Master HR proposed level", all(supl[(r[th["Emp ID"]], r[th["Employee name"]])] == r[hr] for r in tr))
    check("Untouched, the final proposed band equals HR's for everyone", all(r[th["FINAL PROPOSED BAND"]] == r[hr] for r in tr))
    wcols = [i for x, i in th.items() if str(x).endswith("· weightage")]
    filled = sum(1 for r in tr for c in wcols if r[c] in ("L", "M", "H"))
    check("Weightage is filled live from the Department Ladder for every placed person", filled == 7 * sum(1 for r in tr if V.band(r[hr]) != "X"), f"{filled} cells")
    check("Experience and Education are separate parameters", any("Experience · weightage" == x for x in th) and any("Education · weightage" == x for x in th))
    fam = [n for n in bh.sheetnames if n.startswith("RB · ")]
    check("A role & band matrix sheet for every job family", len(fam) >= 19, f"{len(fam)} sheets")
    check("Backend Data is the last sheet", bh.sheetnames[-1] == "Backend Data")
    links = sum(1 for r in tr if r[th["All their data ▸"]] == "Open ▸")
    check("Every row links to the person's raw data", links == 586, f"{links}")
    bh.close()
    live_test()


def live_test():
    """Edit a copy as a business head would, recalculate in LibreOffice, and read the results back."""
    work = tempfile.mkdtemp(prefix="ja_live_")
    wb = openpyxl.load_workbook(BH)
    ws = wb["2. Talent Mapping"]
    hdr = {c.value: c.column for c in ws[6]}
    hrc = hdr["HR proposed band · market benchmark and people inputs"]
    up = next(r for r in range(7, ws.max_row + 1) if ws.cell(r, hrc).value == "B1" and ws.cell(r, hdr["Track"]).value == "IC")
    for c in [v for k, v in hdr.items() if str(k).endswith("· assessment")]:
        ws.cell(up, c).value = "H"
    down = next(r for r in range(up + 1, ws.max_row + 1) if ws.cell(r, hrc).value == "A3")
    ws.cell(down, hdr["Complexity & Expertise · assessment"]).value = "L"
    ws.cell(down, hdr["Autonomy & Decision · assessment"]).value = "L"
    over = next(r for r in range(up + 1, ws.max_row + 1) if ws.cell(r, hrc).value == "A3" and r != down)
    ws.cell(over, hdr["Business head override · optional"]).value = "E2"
    dl = wb["5. Department Ladder"]
    dh = {c.value: c.column for c in dl[6]}
    dl.cell(7, dh["Complexity & Expertise"]).value = "L" if dl.cell(7, dh["Complexity & Expertise"]).value != "L" else "H"
    key = dl.cell(7, dh["Ladder key"]).value
    path = os.path.join(work, "edit.xlsx")
    wb.save(path)
    out = X.recalc_copy(path, work)
    v = openpyxl.load_workbook(out, read_only=True, data_only=True)
    t = list(v["2. Talent Mapping"].iter_rows(min_row=6, values_only=True))
    th = {x: i for i, x in enumerate(t[0])}
    r_up, r_down, r_over = t[up - 6], t[down - 6], t[over - 6]
    check("Live: all seven assessments High move a B1 one level up", r_up[th["FINAL PROPOSED BAND"]] == "B2", str(r_up[th["FINAL PROPOSED BAND"]]))
    check("Live: Complexity and Autonomy Low move an A3 one level down", r_down[th["FINAL PROPOSED BAND"]] == "A2",
          f"{r_down[th['FINAL PROPOSED BAND']]} · net {r_down[th['Weighted net · −1 to +1, live']]}")
    check("Live: an override that breaks the manager rule is flagged", str(r_over[th["Manager rule check · live"]]).startswith("⚠"),
          str(r_over[th["Manager rule check · live"]]))
    ladder_key = th["Ladder key"]
    holders = [r for r in t[1:] if r[ladder_key] == key]
    newv = list(v["5. Department Ladder"].iter_rows(min_row=7, max_row=7, values_only=True))[0][dh["Complexity & Expertise"] - 1]
    check("Live: a Department Ladder weightage edit reaches its holders", holders and all(r[th["Complexity & Expertise · weightage"]] == newv for r in holders),
          f"{len(holders)} holders now {newv}")
    errs = sum(1 for ws in v.worksheets for row in ws.iter_rows(values_only=True) for x in row if isinstance(x, str) and x.startswith("#") and x.endswith(("!", "?", "A")))
    check("Live: no formula errors after the edits", errs == 0, f"{errs}")
    v.close()
    shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
