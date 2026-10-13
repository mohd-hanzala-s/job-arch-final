"""Market evidence for the Role Journey and the References sheet, read from the v1.15 HR Super Master.

R3 in '11. References & Benchmarks' holds, for every department and job family, the IRIS title at each level against
competitors, regtech and SaaS leaders, standards and certification paths. The References sheet reproduces sheet 11 in
full, the level crosswalk and experience triangulation from the Master Framework (M13), the weight benchmark (F21), and
the sources behind the v2.1 five-step scales.
"""
import re

import openpyxl

import ja_common as J

LEVEL_COLS = ["A1", "A2", "A3", "B1", "B2", "B3", "B+", "C1", "C2", "C+", "D1", "D2", "E1", "E2"]

METHOD_SOURCES = [
    ("S70", "Saaty, T. L. (2008) — Decision making with the analytic hierarchy process. International Journal of Services Sciences 1(1)",
     "Peer-reviewed method", "The 1–9 fundamental scale of relative importance; the five weightage steps count 1:3:5:7:9 — its odd steps",
     "https://doi.org/10.1504/ijssci.2008.017590", "Primary"),
    ("S71", "Edwards, W. & Barron, F. H. (1994) — SMARTS and SMARTER: improved simple methods for multiattribute utility measurement. "
     "Organizational Behavior and Human Decision Processes 60(3), 306–325", "Peer-reviewed method",
     "Direct-rating weights: each parameter's share = its importance ÷ the sum of importances", "https://doi.org/10.1006/obhd.1994.1087",
     "Primary"),
    ("S72", "NHS Job Evaluation Handbook — weighting principles", "Point-factor plan",
     "Factor levels carry points, and weighting sets how much each factor contributes; knowledge weighted highest; every job scores at "
     "least one on each factor", "https://www.gov.scot/publications/nhs-job-evaluation-handbook-second-edition/pages/4/", "Primary"),
    ("S73", "European Institute for Gender Equality — gender-neutral job evaluation toolkit, point-factor method",
     "Public-sector guidance", "Levels per sub-factor turned into points by the weight each sub-factor carries",
     "https://eige.europa.eu/gender-mainstreaming/toolkits/gender-neutral-job-evaluation/large-organisations/act/"
     "step-1-understand-point-factor-method-and-factor-and-subfactor-plan", "Primary"),
    ("S74", "ILO / ITCILO — gender-neutral job evaluation for equal pay: factors, sub-factors and weighting",
     "International standard-setter", "Factors, sub-factors and weighting in point-factor job evaluation",
     "https://www.itcilo.org/courses/closing-gender-gap-ilo-approach-gender-neutral-job-evaluation-equal-pay", "Secondary"),
    ("S75", "McKinsey — How to identify the right spans of control for your organization", "Consulting research",
     "Five manager archetypes from player/coach to coordinator; the player/coach carries real individual work. Basis for opening the "
     "B+ and C+ Player-Coach steps to managers who still deliver, and for span norms",
     "https://www.mckinsey.com/capabilities/people-and-organization/our-insights/how-to-identify-the-right-spans-of-control-for-your-organization",
     "Secondary"),
    ("S76", "Deloitte — Spans and layers for the modern organization", "Consulting research",
     "No single right span or number of layers: spans are flagged against norms, never forced",
     "https://www2.deloitte.com/us/en/pages/human-capital/articles/spans-and-layers-for-modern-organization.html", "Secondary"),
    ("S77", "Mercer — International Position Evaluation (IPE): Impact, Communication, Innovation, Knowledge, Risk",
     "Job-evaluation method", "Factor definitions behind the seven IRIS parameters; knowledge includes people-management skill",
     "https://www.mercer.com/en-us/solutions/talent-and-rewards/job-architecture/job-evaluation-ipe", "Primary"),
    ("S78", "Payscale — Job architecture quick guide", "Practitioner guide",
     "Families and sub-families are separate from levels; levels are defined once and used the same way across families — the basis "
     "for a role = sub-family on one shared ladder", "https://www.payscale.com/content/quick-guide/job-architecture-quick-guide.pdf",
     "Secondary"),
    ("S79", "WorldatWork — Dual career ladders: ways to reward specialist leadership", "Professional body",
     "Parallel individual-contributor and management tracks: the IC and Manager lanes inside each role in Bands B and C",
     "https://worldatwork.org/workspan/articles/dual-career-ladders-ways-to-reward-specialist-leadership", "Secondary"),
    ("S80", "Mercer — Job architecture and career frameworks", "Consulting method",
     "Career frameworks set levels, tracks and families / sub-families together; only one in two companies has one",
     "https://www.mercer.com/en-ph/solutions/talent-and-rewards/job-architecture/", "Secondary"),
]


def _rows(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)] if sheet in wb.sheetnames else []
    wb.close()
    return rows


def _cells(r):
    return [x for x in r if x not in (None, "")]


def r3_blocks(path):
    """[{dept, family, standard, rows: [(label, {level: value})]}] from R3."""
    rows = _rows(path, "11. References & Benchmarks")
    out, dept, cur, cols = [], "", None, None
    on = False
    for r in rows:
        c = _cells(r)
        first = J.text(c[0]) if c else ""
        if first.startswith("R3."):
            on = True
            continue
        if first.startswith("R4."):
            break
        if not on or not c:
            continue
        if len(c) == 1 and re.search(r" — \d+ people$", first):
            dept = first.split(" — ")[0]
            continue
        m = re.match(r"^IRIS (.+?) Vs · Standard: (.*)$", first)
        if len(c) == 1 and m:
            cur = {"dept": dept, "family": m.group(1).strip(), "standard": m.group(2).strip(), "rows": []}
            out.append(cur)
            cols = None
            continue
        if first == "Ladder / company":
            cols = {J.text(x): j for j, x in enumerate(r) if J.text(x) in LEVEL_COLS}
            continue
        if cur is not None and cols:
            lab = J.text(r[0]) if r[0] not in (None, "") else first
            vals = {lv: J.text(r[j]) for lv, j in cols.items() if j < len(r) and J.has(r[j])}
            if vals:
                cur["rows"].append((lab, vals))
    return out


def market_rows(blocks, family, dept, manager):
    """The market rows for a role: its family's block in its department if there is one, else the family's first block.
    Benchmark composites first (managers' row first for a manager role), then named companies; IRIS's own rows left out."""
    cand = [b for b in blocks if b["family"] == family]
    if not cand:
        return None, []
    b = next((x for x in cand if x["dept"] == dept), cand[0])
    rows = [(l, v) for l, v in b["rows"] if not l.startswith("IRIS —")]
    bench = [(l, v) for l, v in rows if l.startswith("IRIS Vs")]
    comp = [(l, v) for l, v in rows if not l.startswith("IRIS Vs")]
    if manager:
        bench.sort(key=lambda x: 0 if "manager" in x[0].lower() else 1)
    else:
        bench.sort(key=lambda x: 1 if "manager" in x[0].lower() else 0)
    return b, bench + comp


def short_label(l):
    l = re.sub(r"^IRIS Vs (Competitor|Market) Benchmark — ", "", l)
    return re.split(r" · | — ", l)[0][:40]


def reference_sections(path):
    """[(title, rows)] — sheet 11 in full, M13 from the Master Framework, and the weight benchmark."""
    r11 = [[("Department not recorded" + J.text(x)[len("Not provided · HRIS"):] if isinstance(x, str) and x.startswith("Not provided · HRIS")
             else ("IRIS Unassigned Vs · Standard: none" + J.text(x)[len("IRIS Unassigned Vs · Standard: —"):]
                   if isinstance(x, str) and x.startswith("IRIS Unassigned Vs · Standard: —") else x)) for x in r]
           for r in _rows(path, "11. References & Benchmarks")[4:]]
    out = [("Department × level benchmarks, peer universe, professional standards and the source register · carried over in full from the HR Super Master", r11)]
    fw = _rows(path, "F14. Master Framework")
    a = next((i for i, r in enumerate(fw) if _cells(r) and J.text(_cells(r)[0]).startswith("M13.")), None)
    b = next((i for i, r in enumerate(fw) if _cells(r) and J.text(_cells(r)[0]).startswith("M14.")), None)
    if a is not None and b is not None:
        out.append(("M13 · Market benchmarking — methods, level crosswalk, experience triangulation · Master Framework", fw[a:b]))
    wt = _rows(path, "F21. Weight Benchmark")
    if wt:
        out.append(("Weight benchmark — factor weights by role group and source · F21", wt[4:]))
    return out
